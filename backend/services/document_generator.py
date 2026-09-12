"""
ResumeIQ — ATS-Safe Document Generation Engine
Renders clean, ATS-safe PDF (via ReportLab) and DOCX (via python-docx) resume files.
"""

import os
import logging
from pathlib import Path
from typing import Dict, Any, Optional

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT

from backend.config import get_settings

logger = logging.getLogger(__name__)


class DocumentGeneratorService:
    """Generates clean ATS-parseable PDF and DOCX documents."""

    def generate_pdf(self, structured_content: Dict[str, Any], output_filename: str) -> str:
        """Render ATS-safe PDF using ReportLab."""
        settings = get_settings()
        out_dir = Path(settings.UPLOAD_DIR) / "generated"
        out_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = out_dir / output_filename

        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            alignment=TA_CENTER,
            textColor='#1a202c',
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'DocSubTitle',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            alignment=TA_CENTER,
            textColor='#4a5568',
            spaceAfter=12
        )
        heading_style = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            alignment=TA_LEFT,
            textColor='#2b6cb0',
            spaceBefore=10,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            alignment=TA_LEFT,
            textColor='#2d3748',
            spaceAfter=4
        )

        story = []

        # Header
        title_text = structured_content.get("name", "Candidate Resume")
        story.append(Paragraph(f"<b>{title_text}</b>", title_style))
        role_text = structured_content.get("title", "Target Role")
        story.append(Paragraph(f"<b>{role_text}</b>", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1, color='#cbd5e0', spaceAfter=10))

        # Summary
        if "summary" in structured_content:
            story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style))
            story.append(Paragraph(structured_content["summary"], body_style))
            story.append(Spacer(1, 8))

        # Skills
        if "verified_skills" in structured_content:
            story.append(Paragraph("<b>TECHNICAL SKILLS</b>", heading_style))
            skills_str = ", ".join(structured_content["verified_skills"])
            story.append(Paragraph(f"<b>Core Competencies:</b> {skills_str}", body_style))
            story.append(Spacer(1, 8))

        # Experience Bullets
        if "bullets" in structured_content:
            story.append(Paragraph("<b>PROFESSIONAL EXPERIENCE & HIGHLIGHTS</b>", heading_style))
            for bullet_obj in structured_content["bullets"]:
                b_text = bullet_obj["bullet"] if isinstance(bullet_obj, dict) else str(bullet_obj)
                story.append(Paragraph(f"• {b_text}", body_style))
            story.append(Spacer(1, 8))

        doc.build(story)
        logger.info(f"Generated ATS-safe PDF at {pdf_path}")
        return str(pdf_path)

    def generate_docx(self, structured_content: Dict[str, Any], output_filename: str) -> str:
        """Render ATS-safe DOCX using python-docx."""
        try:
            from docx import Document
        except ImportError:
            logger.warning("python-docx not installed")
            return ""

        settings = get_settings()
        out_dir = Path(settings.UPLOAD_DIR) / "generated"
        out_dir.mkdir(parents=True, exist_ok=True)
        docx_path = out_dir / output_filename

        doc = Document()
        doc.add_heading(structured_content.get("name", "Candidate Resume"), level=0)
        doc.add_paragraph(structured_content.get("title", "Target Role"))

        if "summary" in structured_content:
            doc.add_heading("PROFESSIONAL SUMMARY", level=1)
            doc.add_paragraph(structured_content["summary"])

        if "verified_skills" in structured_content:
            doc.add_heading("TECHNICAL SKILLS", level=1)
            doc.add_paragraph(", ".join(structured_content["verified_skills"]))

        if "bullets" in structured_content:
            doc.add_heading("EXPERIENCE HIGHLIGHTS", level=1)
            for bullet_obj in structured_content["bullets"]:
                b_text = bullet_obj["bullet"] if isinstance(bullet_obj, dict) else str(bullet_obj)
                doc.add_paragraph(b_text, style='List Bullet')

        doc.save(str(docx_path))
        logger.info(f"Generated ATS-safe DOCX at {docx_path}")
        return str(docx_path)


document_generator_service = DocumentGeneratorService()
